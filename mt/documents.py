"""Meeting documents: files you add while recording, to give the summary more context.

Layout inside a meeting folder:
    documents/<name>              the copy of the file you chose (a second file with the same name becomes "name (2).ext")
    documents/extracted/<name>.txt  the text read from it (with "--- page 3 ---" style markers)
    documents/documents.json      what was added and when (seconds into the recording), and how reading is going

Claude is not given the text. It gets a list with a short preview of each document and opens the extracted text files
that matter, by path (see summarize.py). Reading happens in the background, one file at a time (DocumentWorker); the
meeting's own summary waits for it.
"""

import logging
import hashlib
import json
import queue
import re
import shutil
import threading
import uuid
import zipfile
from pathlib import Path

from . import ocr

log = logging.getLogger("mt.documents")

MAX_DOCUMENTS = 30
MAX_BYTES = 100 * 1024 * 1024
MAX_PAGES = 300  # pages (PDF) or pictures (everything else) that are read for text in images
PDF = {".pdf"}
WORD = {".docx"}
SLIDES = {".pptx"}
SHEET = {".xlsx"}
TEXT = {".txt", ".md", ".markdown"}
IMAGE = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".heic", ".bmp", ".gif"}
SUPPORTED = PDF | WORD | SLIDES | SHEET | TEXT | IMAGE
KIND = {**{e: "pdf" for e in PDF}, **{e: "docx" for e in WORD}, **{e: "pptx" for e in SLIDES}, **{e: "xlsx" for e in SHEET},
        **{e: "txt" for e in TEXT}, **{e: "img" for e in IMAGE}}
KIND_LABEL = {"pdf": "PDF", "docx": "Word", "pptx": "PowerPoint", "xlsx": "Excel", "txt": "Text", "img": "Image"}
PREVIEW_CHARS = 300

_lock = threading.RLock()  # the manifest is touched by the app and by the reading thread


class Cancelled(Exception):
    pass


# ---- paths and the manifest
def docs_dir(meeting_dir: Path) -> Path:
    return Path(meeting_dir) / "documents"


def extracted_dir(meeting_dir: Path) -> Path:
    return docs_dir(meeting_dir) / "extracted"


def _manifest(meeting_dir: Path) -> Path:
    return docs_dir(meeting_dir) / "documents.json"


def load(meeting_dir: Path) -> list[dict]:
    """The documents of a meeting. Entries whose file was removed by hand are dropped."""
    with _lock:
        try:
            items = json.loads(_manifest(meeting_dir).read_text()).get("documents", [])
        except (OSError, ValueError):
            return []
        return [d for d in items if isinstance(d, dict)
                and (d.get("status") == "queued" or (docs_dir(meeting_dir) / d.get("stored", "")).exists())]  # (queued: not copied yet)


def _save(meeting_dir: Path, items: list[dict]) -> None:
    with _lock:
        d = docs_dir(meeting_dir)
        d.mkdir(parents=True, exist_ok=True)
        tmp = _manifest(meeting_dir).with_suffix(".json.tmp")
        tmp.write_text(json.dumps({"documents": items}, indent=2, ensure_ascii=False))
        tmp.replace(_manifest(meeting_dir))


def update(meeting_dir: Path, doc_id: str, **fields) -> bool:
    """Change fields of one entry. False if it no longer exists (it was deleted)."""
    with _lock:
        items = load(meeting_dir)
        for d in items:
            if d["id"] == doc_id:
                d.update(fields)
                _save(meeting_dir, items)
                return True
        return False


def get(meeting_dir: Path, doc_id: str) -> dict | None:
    return next((d for d in load(meeting_dir) if d["id"] == doc_id), None)


def count(meeting_dir: Path | None) -> int:
    return len(load(meeting_dir)) if meeting_dir else 0


def _sha(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def add(meeting_dir: Path, sources: list[Path], added_at: int, ocr_on: bool) -> tuple[list[dict], list[str]]:
    """Register files for a meeting (quick: nothing is copied yet). -> (new entries, messages about skipped files).
    The copying and reading happen in the DocumentWorker."""
    new, skipped = [], []
    with _lock:
        items = load(meeting_dir)
        known = {d.get("sha") for d in items if d.get("sha")}
        for src in sources:
            src = Path(src)
            if src.suffix.lower() not in SUPPORTED:
                skipped.append(f"{src.name}: this kind of file is not supported")
                continue
            if not src.is_file():
                skipped.append(f"{src.name}: not found")
                continue
            if src.stat().st_size > MAX_BYTES:
                skipped.append(f"{src.name}: bigger than {MAX_BYTES // (1024 * 1024)} MB")
                continue
            if len(items) + len(new) >= MAX_DOCUMENTS:
                skipped.append(f"{src.name}: a meeting can have {MAX_DOCUMENTS} documents")
                continue
            digest = _sha(src)
            if digest in known:
                skipped.append(f"{src.name}: already added")
                continue
            known.add(digest)
            taken = {d["stored"] for d in items + new}
            stored = src.name
            n = 2
            while stored in taken or (docs_dir(meeting_dir) / stored).exists():
                stored = f"{src.stem} ({n}){src.suffix}"
                n += 1
            kind = KIND[src.suffix.lower()]
            new.append({"id": uuid.uuid4().hex[:12], "name": src.name, "stored": stored, "kind": kind,
                        "source": str(src), "sha": digest, "added_at": int(added_at), "size": src.stat().st_size,
                        "ocr": bool(ocr_on), "status": "queued", "progress": 0.0, "detail": "", "note": "", "chars": 0})
        if new:
            _save(meeting_dir, items + new)
    return new, skipped


def delete(meeting_dir: Path, doc_id: str) -> bool:
    """Remove the copy, the text read from it and its entry (its timestamp goes with it). The original is never touched."""
    with _lock:
        items = load(meeting_dir)
        gone = next((d for d in items if d["id"] == doc_id), None)
        if gone is None:
            return False
        _save(meeting_dir, [d for d in items if d["id"] != doc_id])
        if gone.get("status") in ("queued", "reading"):  # only a file that is still being read needs stopping
            _cancelled.add(doc_id)
    _remove_files(meeting_dir, gone)
    return True


def _remove_files(meeting_dir: Path, d: dict) -> None:
    (docs_dir(meeting_dir) / d["stored"]).unlink(missing_ok=True)
    (extracted_dir(meeting_dir) / f"{d['stored']}.txt").unlink(missing_ok=True)


def extracted_file(meeting_dir: Path, d: dict) -> Path:
    return extracted_dir(meeting_dir) / f"{d['stored']}.txt"


def preview(text: str) -> str:
    t = re.sub(r"---[^-\n]*---", " ", text)  # the page markers
    t = re.sub(r"\s+", " ", t).strip()
    return (t[:PREVIEW_CHARS].rstrip() + "…") if len(t) > PREVIEW_CHARS else t


# ---- reading the text out of each kind of file
_cancelled: set[str] = set()


def _check(doc_id: str) -> None:
    if doc_id in _cancelled:
        raise Cancelled()


def _read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-16"):
        try:
            return raw.decode(enc)
        except UnicodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _media_images(path: Path, prefix: str) -> list[tuple[str, bytes]]:
    """The pictures inside a Word / PowerPoint / Excel file (they are zip files), without duplicates, in file order."""
    out, seen = [], set()
    with zipfile.ZipFile(path) as z:
        for name in sorted(z.namelist(), key=lambda n: [int(x) if x.isdigit() else x for x in re.split(r"(\d+)", n)]):
            if name.startswith(prefix) and Path(name).suffix.lower() in IMAGE:
                data = z.read(name)
                h = hashlib.sha1(data).hexdigest()
                if h not in seen:
                    seen.add(h)
                    out.append((name, data))
    return out


def _ocr_pictures(pictures, progress, doc_id) -> list[str]:
    chunks = []
    pictures = pictures[:MAX_PAGES]
    for i, (name, data) in enumerate(pictures, 1):
        _check(doc_id)
        try:
            cg = ocr.image_from_bytes(data)
            if cg is not None and ocr.is_worth_reading(cg):
                lines = ocr.read_image(cg)
                if lines:
                    chunks.append(f"--- picture {i} ({Path(name).name}) ---\n" + "\n".join(lines))
        except Exception:
            log.exception("OCR of an embedded picture failed")
        progress(i / len(pictures))
    return chunks


def read_pdf(path: Path, ocr_on: bool, progress, doc_id: str) -> tuple[str, str, str]:
    """-> (text, detail, note). Page text where the PDF has it; with OCR on, every page is also read as a picture and the
    lines that are not in the page's text are added (screenshots, scans)."""
    from Foundation import NSURL
    from Quartz import PDFDocument

    doc = PDFDocument.alloc().initWithURL_(NSURL.fileURLWithPath_(str(path)))
    if doc is None:
        raise RuntimeError("not a readable PDF")
    if doc.isLocked():
        raise RuntimeError("Password-protected")
    pages = int(doc.pageCount())
    parts, scanned = [], 0
    can_ocr = ocr_on and ocr.available()
    for i in range(pages):
        _check(doc_id)
        page = doc.pageAtIndex_(i)
        layer = str(page.string() or "").strip()
        text = layer
        if can_ocr and i < MAX_PAGES:
            try:
                lines = ocr.read_image(_render_page(page))
                extra = ocr.new_lines(lines, layer) if layer else lines
                if extra:
                    text = (layer + "\n" if layer else "") + "\n".join(extra)
                    if not layer:
                        scanned += 1
            except Exception:
                log.exception("OCR of a page failed")
        if text:
            parts.append(f"--- page {i + 1} ---\n{text}")
        progress((i + 1) / max(pages, 1))
    detail = f"{pages} page{'s' if pages != 1 else ''}"
    note = ""
    if not parts:
        note = "Scanned pages: turn on OCR to read them" if not ocr_on else ""
    elif pages > MAX_PAGES and can_ocr:
        note = f"Text in images was read on the first {MAX_PAGES} pages only"
    return "\n\n".join(parts), detail, note


def _render_page(page, dpi: int = 150):
    """A PDF page as a picture (150 dpi: enough to read 6 pt text, see ocr.py)."""
    import Quartz

    box = page.boundsForBox_(Quartz.kPDFDisplayBoxMediaBox)
    w, h = max(1, int(box.size.width * dpi / 72)), max(1, int(box.size.height * dpi / 72))
    cs = Quartz.CGColorSpaceCreateDeviceRGB()
    ctx = Quartz.CGBitmapContextCreate(None, w, h, 8, 0, cs, Quartz.kCGImageAlphaPremultipliedLast)
    Quartz.CGContextSetRGBFillColor(ctx, 1, 1, 1, 1)
    Quartz.CGContextFillRect(ctx, ((0, 0), (w, h)))
    Quartz.CGContextScaleCTM(ctx, dpi / 72, dpi / 72)
    page.drawWithBox_toContext_(Quartz.kPDFDisplayBoxMediaBox, ctx)
    return Quartz.CGBitmapContextCreateImage(ctx)


def read_docx(path: Path, ocr_on: bool, progress, doc_id: str) -> tuple[str, str, str]:
    import pypandoc

    text = pypandoc.convert_file(str(path), "plain", format="docx", extra_args=["--wrap=none"]).strip()
    pictures = _media_images(path, "word/media/")
    detail = f"{len(pictures)} picture{'s' if len(pictures) != 1 else ''}" if pictures else "Text"
    chunks = _ocr_pictures(pictures, progress, doc_id) if (ocr_on and pictures and ocr.available()) else []
    return "\n\n".join([t for t in [text] + chunks if t]), detail, ""


def read_pptx(path: Path, ocr_on: bool, progress, doc_id: str) -> tuple[str, str, str]:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        _check(doc_id)
        texts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    texts.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                texts.append("Notes: " + notes)
        if texts:
            parts.append(f"--- slide {i} ---\n" + "\n".join(texts))
    pictures = _media_images(path, "ppt/media/")
    chunks = _ocr_pictures(pictures, progress, doc_id) if (ocr_on and pictures and ocr.available()) else []
    n = len(prs.slides)
    return "\n\n".join(parts + chunks), f"{n} slide{'s' if n != 1 else ''}", ""


def read_xlsx(path: Path, ocr_on: bool, progress, doc_id: str) -> tuple[str, str, str]:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        _check(doc_id)
        rows = []
        for n, row in enumerate(ws.iter_rows(values_only=True)):
            if n >= 2000:
                rows.append("(more rows not read)")
                break
            cells = ["" if v is None else str(v) for v in row]
            if any(cells):
                rows.append(" | ".join(cells).rstrip(" |"))
        if rows:
            parts.append(f"--- sheet {ws.title} ---\n" + "\n".join(rows))
    sheets = len(wb.worksheets)
    wb.close()
    pictures = _media_images(path, "xl/media/")
    chunks = _ocr_pictures(pictures, progress, doc_id) if (ocr_on and pictures and ocr.available()) else []
    return "\n\n".join(parts + chunks), f"{sheets} sheet{'s' if sheets != 1 else ''}", ""


def read_image(path: Path, ocr_on: bool, progress, doc_id: str) -> tuple[str, str, str]:
    if not ocr_on:
        return "", "Picture", "Turn on OCR to read text in pictures"
    if not ocr.available():
        return "", "Picture", "Text recognition is not available on this Mac"
    cg = ocr.image_from_bytes(path.read_bytes())
    if cg is None:
        raise RuntimeError("not a readable picture")
    lines = ocr.read_image(cg)
    progress(1.0)
    return "\n".join(lines), "Picture", ""


READERS = {"pdf": read_pdf, "docx": read_docx, "pptx": read_pptx, "xlsx": read_xlsx, "img": read_image}


def process(meeting_dir: Path, doc: dict, on_change=lambda: None) -> None:
    """Copy one registered file into the meeting folder and read its text. Runs in the DocumentWorker."""
    doc_id = doc["id"]
    folder = docs_dir(meeting_dir)

    def set_(**f):
        if not update(meeting_dir, doc_id, **f):
            raise Cancelled()
        on_change()

    try:
        _check(doc_id)
        folder.mkdir(parents=True, exist_ok=True)
        dst = folder / doc["stored"]
        if not dst.exists():
            tmp = dst.with_name(dst.name + ".part")
            shutil.copyfile(doc["source"], tmp)
            _check(doc_id)
            tmp.replace(dst)
        set_(status="reading", progress=0.0)
        last = [0.0]

        def progress(f):
            if f - last[0] >= 0.02 or f >= 1:
                last[0] = f
                set_(progress=round(f, 3))

        if doc["kind"] == "txt":
            text, detail, note = _read_text_file(dst), "Text", ""
        else:
            text, detail, note = READERS[doc["kind"]](dst, doc["ocr"], progress, doc_id)
        _check(doc_id)
        text = text.strip()
        extracted_dir(meeting_dir).mkdir(parents=True, exist_ok=True)
        if text:
            extracted_file(meeting_dir, doc).write_text(text + "\n", encoding="utf-8")
        set_(status="ready" if text else "nolext", progress=1.0, detail=detail, note=note, chars=len(text),
             preview=preview(text) if text else "")
    except Cancelled:
        pass  # deleted while it was being read: delete() already removed what it knew about
    except Exception as e:
        log.exception("Reading document %s failed", doc.get("name"))
        try:
            update(meeting_dir, doc_id, status="failed", note=str(e)[:80] or "Couldn't read this file")
            on_change()
        except Exception:
            pass
    finally:
        if doc_id in _cancelled:  # clean up whatever the reader wrote after the delete
            _remove_files(meeting_dir, doc)
        _cancelled.discard(doc_id)


class DocumentWorker:
    """Reads documents one at a time, at low priority, so recording and transcription stay smooth."""

    def __init__(self, on_change=lambda: None):
        self.on_change = on_change
        self._q: queue.Queue = queue.Queue()
        self._idle = threading.Condition()
        self._busy = 0
        threading.Thread(target=self._run, daemon=True).start()

    def submit(self, meeting_dir: Path, docs: list[dict]) -> None:
        with self._idle:
            self._busy += len(docs)
        for d in docs:
            self._q.put((Path(meeting_dir), d))
        self.on_change()

    def pending(self, meeting_dir: Path | None = None) -> int:
        """How many documents of this meeting are still waiting or being read."""
        if meeting_dir is None:
            return self._busy
        return sum(1 for d in load(meeting_dir) if d.get("status") in ("queued", "reading"))

    def wait(self, meeting_dir: Path, on_progress=lambda done, total: None, timeout: float = 3600) -> None:
        """Block until this meeting's documents are read (or `timeout`), reporting 'n of m'."""
        import time

        deadline = time.time() + timeout
        total = max(count(meeting_dir), 1)
        while time.time() < deadline:
            left = self.pending(meeting_dir)
            if left == 0:
                return
            on_progress(total - left, total)
            time.sleep(0.5)

    def _run(self) -> None:
        try:
            import ctypes

            ctypes.CDLL("/usr/lib/libSystem.B.dylib").pthread_set_qos_class_self_np(0x11, 0)  # QOS_CLASS_UTILITY
        except Exception:
            pass
        while True:
            meeting_dir, doc = self._q.get()
            try:
                process(meeting_dir, doc, self.on_change)
            finally:
                with self._idle:
                    self._busy -= 1
                self.on_change()


def label(d: dict) -> str:
    return KIND_LABEL.get(d.get("kind", ""), "File")
